#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
P0 inventory and requirement extraction for the local course project package.

Python 3.10+.
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import sys
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


DEFAULT_KEYWORDS = [
    "变量选择",
    "前向分步回归",
    "熵权-PROMETHEE",
    "下一期收益率匹配",
    "评级",
    "回测指标",
]

DEFAULT_ASSUMPTIONS = [
    ("business_goal", "用多因子与宏观因子辅助变量选择、评级与收益预测", "default_assumption", "业务目标需由作业要求和建模目标确认"),
    ("label_definition", "下一期收益率", "default_assumption", "标签定义需确认收益率口径、频率和是否复权"),
    ("macro_alignment_mode", "按预测时点可获得的上一期宏观数据对齐", "default_assumption", "防止未来宏观数据泄露"),
    ("rating_bins", "5", "default_assumption", "默认五档评级，可按业务解释性调整"),
    ("transaction_cost", "0.001", "default_assumption", "默认单边交易成本 10bp，需按回测要求确认"),
    ("rebalancing_freq", "quarterly", "default_assumption", "默认按季度调仓，与季度 sheet 粒度一致"),
    ("validation_scheme", "time_series_split", "default_assumption", "金融预测默认使用时间序列验证"),
    ("industry_neutralization", "false", "default_assumption", "是否行业中性化需结合数据字段确认"),
    ("leakage_checks", "enabled", "default_assumption", "默认启用标签期、宏观期、标准化拟合窗口检查"),
    ("stock_universe", "all_available_stocks", "default_assumption", "股票池范围需确认是否剔除 ST、停牌、上市不足样本"),
    ("missing_value_strategy", "median_by_period", "default_assumption", "默认按期中位数填补，需记录缺失比例"),
    ("outlier_treatment", "winsorize_by_period_1pct_99pct", "default_assumption", "默认按期缩尾以降低极端值影响"),
    ("factor_standardization", "zscore_by_period", "default_assumption", "横截面因子默认按期标准化"),
    ("benchmark", "equal_weight_universe", "default_assumption", "默认与全样本等权基准比较"),
    ("promethee_preference_function", "usual", "default_assumption", "PROMETHEE 偏好函数需结合指标方向和阈值确认"),
    ("forward_stepwise_criterion", "aic", "default_assumption", "前向分步回归默认用 AIC，也可改为 BIC、调整 R2 或验证集指标"),
    ("random_seed", "2026", "default_assumption", "保证可复现实验"),
]


@dataclass
class InventoryFile:
    relative_path: str
    size_bytes: int
    suffix: str


def readable_error(message: str) -> SystemExit:
    return SystemExit(f"[ERROR] {message}")


def contains_cjk(text: str) -> bool:
    return any("\u4e00" <= char <= "\u9fff" for char in text)


def normalize_zip_member_name(name: str) -> str:
    """Repair common Windows ZIP Chinese filename mojibake when UTF-8 flag is absent."""
    if contains_cjk(name):
        return name
    try:
        repaired = name.encode("cp437").decode("gbk")
    except UnicodeError:
        return name
    return repaired if contains_cjk(repaired) else name


def resolve_existing_file(path_text: str, base_dir: Path, label: str) -> Path:
    path = Path(path_text)
    if not path.is_absolute():
        path = base_dir / path
    path = path.resolve()
    if not path.exists():
        raise readable_error(f"{label} 不存在: {path}")
    if not path.is_file():
        raise readable_error(f"{label} 不是文件: {path}")
    return path


def safe_extract_zip(zip_path: Path, extract_dir: Path) -> list[str]:
    extract_dir.mkdir(parents=True, exist_ok=True)
    extracted_members: list[str] = []
    extract_root = extract_dir.resolve()
    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            for member in zf.infolist():
                member_name = normalize_zip_member_name(member.filename).replace("\\", "/")
                if Path(member_name).is_absolute() or ".." in Path(member_name).parts:
                    raise readable_error(f"压缩包包含不安全路径，已拒绝解压: {member.filename}")
                target_path = (extract_root / member_name).resolve()
                try:
                    target_path.relative_to(extract_root)
                except ValueError as exc:
                    raise readable_error(f"压缩包包含越界路径，已拒绝解压: {member.filename}") from exc

                extracted_members.append(member_name)
                if member.is_dir() or member_name.endswith("/"):
                    target_path.mkdir(parents=True, exist_ok=True)
                    continue
                target_path.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(member, "r") as source, target_path.open("wb") as target:
                    shutil.copyfileobj(source, target)
    except zipfile.BadZipFile as exc:
        raise readable_error(f"无法读取 zip 文件，文件可能已损坏: {zip_path}") from exc
    return extracted_members


def inventory_files(root: Path) -> list[InventoryFile]:
    files: list[InventoryFile] = []
    for path in sorted(root.rglob("*")):
        if path.is_file():
            files.append(
                InventoryFile(
                    relative_path=path.relative_to(root).as_posix(),
                    size_bytes=path.stat().st_size,
                    suffix=path.suffix.lower(),
                )
            )
    return files


def locate_extracted_file(extract_dir: Path, relative_text: str, label: str) -> Path:
    candidate = (extract_dir / relative_text).resolve()
    if candidate.exists() and candidate.is_file():
        return candidate

    basename_matches = [path for path in extract_dir.rglob(Path(relative_text).name) if path.is_file()]
    if len(basename_matches) == 1:
        return basename_matches[0].resolve()
    if len(basename_matches) > 1:
        matches = "\n".join(str(path.relative_to(extract_dir)) for path in basename_matches)
        raise readable_error(f"{label} 存在多个同名候选文件，请传入更精确的相对路径:\n{matches}")
    raise readable_error(f"{label} 在解压目录中不存在: {relative_text}")


def read_docx_text(docx_path: Path) -> str:
    try:
        from docx import Document

        document = Document(str(docx_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_cells: list[str] = []
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        table_cells.append(text)
        return "\n".join(paragraphs + table_cells)
    except ImportError:
        try:
            import docx2txt
        except ImportError as exc:
            raise readable_error("读取 docx 需要安装 python-docx 或 docx2txt。建议执行: pip install python-docx") from exc
        return docx2txt.process(str(docx_path)) or ""
    except Exception as exc:
        raise readable_error(f"读取 docx 失败: {docx_path}；原因: {exc}") from exc


def extract_keywords(text: str, keywords: list[str]) -> dict[str, dict[str, Any]]:
    results: dict[str, dict[str, Any]] = {}
    normalized_text = text.replace(" ", "")
    for keyword in keywords:
        normalized_keyword = keyword.replace(" ", "")
        found = normalized_keyword in normalized_text
        snippets: list[str] = []
        if found:
            index = normalized_text.find(normalized_keyword)
            start = max(index - 40, 0)
            end = min(index + len(normalized_keyword) + 40, len(normalized_text))
            snippets.append(normalized_text[start:end])
        results[keyword] = {"found": found, "snippets": snippets}
    return results


def read_excel_sheets(workbook_path: Path, label: str) -> list[str]:
    try:
        import pandas as pd

        excel_file = pd.ExcelFile(workbook_path, engine="openpyxl")
        return list(excel_file.sheet_names)
    except ImportError as exc:
        raise readable_error("读取 xlsx 需要安装 pandas 和 openpyxl。建议执行: pip install pandas openpyxl") from exc
    except Exception as exc:
        raise readable_error(f"读取 {label} sheet 列表失败: {workbook_path}；原因: {exc}") from exc


def validate_expectations(
    inventory: list[InventoryFile],
    stock_sheets: list[str],
    macro_sheets: list[str],
    expected_stock_sheet_count: int,
    expected_macro_sheets: list[str],
) -> list[dict[str, Any]]:
    docx_count = sum(1 for item in inventory if item.suffix == ".docx")
    xlsx_count = sum(1 for item in inventory if item.suffix == ".xlsx")
    checks = [
        {"name": "docx_count", "expected": 1, "actual": docx_count, "passed": docx_count == 1},
        {"name": "xlsx_count", "expected": 2, "actual": xlsx_count, "passed": xlsx_count == 2},
        {
            "name": "stock_sheet_count",
            "expected": expected_stock_sheet_count,
            "actual": len(stock_sheets),
            "passed": len(stock_sheets) == expected_stock_sheet_count,
        },
    ]
    missing_macro_sheets = [sheet for sheet in expected_macro_sheets if sheet not in macro_sheets]
    checks.append(
        {
            "name": "macro_required_sheets",
            "expected": expected_macro_sheets,
            "actual": macro_sheets,
            "missing": missing_macro_sheets,
            "passed": not missing_macro_sheets,
        }
    )
    checks.append(
        {
            "name": "assumption_ledger_config_count",
            "expected": ">=15",
            "actual": len(DEFAULT_ASSUMPTIONS),
            "passed": len(DEFAULT_ASSUMPTIONS) >= 15,
        }
    )
    return checks


def write_inventory_json(
    output_path: Path,
    args: argparse.Namespace,
    extract_dir: Path,
    inventory: list[InventoryFile],
    docx_path: Path,
    stock_path: Path,
    macro_path: Path,
    keyword_results: dict[str, dict[str, Any]],
    stock_sheets: list[str],
    macro_sheets: list[str],
    checks: list[dict[str, Any]],
) -> None:
    payload = {
        "metadata": {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "script": Path(__file__).as_posix(),
            "python_version": sys.version.split()[0],
        },
        "inputs": {
            "zip_path": str(Path(args.zip_path)),
            "task_doc": args.task_doc,
            "stock_workbook": args.stock_workbook,
            "macro_workbook": args.macro_workbook,
            "output_dir": args.output_dir,
        },
        "facts_from_files": {
            "extract_dir": str(extract_dir),
            "recognized_documents": {"docx": docx_path.relative_to(extract_dir).as_posix()},
            "recognized_workbooks": {
                "stock_workbook": stock_path.relative_to(extract_dir).as_posix(),
                "macro_workbook": macro_path.relative_to(extract_dir).as_posix(),
            },
            "files": [asdict(item) for item in inventory],
            "task_keyword_matches": keyword_results,
            "workbook_sheets": {
                "stock_workbook": stock_sheets,
                "macro_workbook": macro_sheets,
            },
            "validation_checks": checks,
        },
        "default_assumptions": [
            {"item": item, "default_value": value, "reason": reason}
            for item, value, _, reason in DEFAULT_ASSUMPTIONS
        ],
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_requirements_md(
    output_path: Path,
    keyword_results: dict[str, dict[str, Any]],
    stock_sheets: list[str],
    macro_sheets: list[str],
    checks: list[dict[str, Any]],
) -> None:
    lines = [
        "# Task Requirements Seed",
        "",
        "## 从文件中识别到的事实",
        "",
        "### 任务关键词",
    ]
    for keyword, result in keyword_results.items():
        status = "识别到" if result["found"] else "未识别到"
        lines.append(f"- {keyword}: {status}")
        for snippet in result["snippets"]:
            lines.append(f"  - 片段: `{snippet}`")

    lines.extend(
        [
            "",
            "### 工作簿 Sheet",
            "",
            f"- 个股工作簿 sheet 数: {len(stock_sheets)}",
            f"- 个股工作簿 sheet 列表: {', '.join(stock_sheets)}",
            f"- 宏观工作簿 sheet 数: {len(macro_sheets)}",
            f"- 宏观工作簿 sheet 列表: {', '.join(macro_sheets)}",
            "",
            "### 验收检查",
        ]
    )
    for check in checks:
        status = "通过" if check["passed"] else "未通过"
        lines.append(f"- {check['name']}: {status}；expected={check['expected']}；actual={check['actual']}")
        if check.get("missing"):
            lines.append(f"  - 缺失: {', '.join(check['missing'])}")

    lines.extend(["", "## 默认假设", ""])
    for item, value, _, reason in DEFAULT_ASSUMPTIONS:
        lines.append(f"- {item}: `{value}`；{reason}")

    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_assumption_ledger_csv(output_path: Path) -> None:
    fieldnames = ["item", "value", "source_type", "status", "configurable", "rationale"]
    with output_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for item, value, source_type, reason in DEFAULT_ASSUMPTIONS:
            writer.writerow(
                {
                    "item": item,
                    "value": value,
                    "source_type": source_type,
                    "status": "to_confirm",
                    "configurable": "true",
                    "rationale": reason,
                }
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="解压本地课程项目包，生成文件清单、任务要求摘要和 assumption ledger 种子文件。"
    )
    parser.add_argument("--zip-path", default="大数据决策-期末大作业.zip", help="待解压的 zip 文件路径")
    parser.add_argument("--task-doc", default="大作业二_期末大作业_变量选择与预测应用.docx", help="任务 docx 在解压目录中的相对路径")
    parser.add_argument(
        "--stock-workbook",
        default="原始数据/因子汇总+月平均收益率(各年份独立)24个因子.xlsx",
        help="个股因子工作簿在解压目录中的相对路径",
    )
    parser.add_argument("--macro-workbook", default="原始数据/宏观因子数据.xlsx", help="宏观因子工作簿在解压目录中的相对路径")
    parser.add_argument("--output-dir", default="outputs", help="输出目录")
    parser.add_argument("--extract-dir", default=None, help="解压目录；默认使用 output_dir/extracted_project")
    parser.add_argument("--expected-stock-sheet-count", type=int, default=30, help="个股工作簿期望 sheet 数")
    parser.add_argument(
        "--expected-macro-sheets",
        default="宏观因子,宏观因子具体数据,Sheet1",
        help="宏观工作簿必须存在的 sheet，逗号分隔",
    )
    parser.add_argument("--fail-on-check", action="store_true", help="验收检查不通过时以非 0 状态退出")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    base_dir = Path.cwd()
    zip_path = resolve_existing_file(args.zip_path, base_dir, "zip_path")
    output_dir = (base_dir / args.output_dir).resolve() if not Path(args.output_dir).is_absolute() else Path(args.output_dir).resolve()
    extract_dir = (
        (base_dir / args.extract_dir).resolve()
        if args.extract_dir and not Path(args.extract_dir).is_absolute()
        else Path(args.extract_dir).resolve()
        if args.extract_dir
        else output_dir / "extracted_project"
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    safe_extract_zip(zip_path, extract_dir)
    inventory = inventory_files(extract_dir)
    docx_path = locate_extracted_file(extract_dir, args.task_doc, "task_doc")
    stock_path = locate_extracted_file(extract_dir, args.stock_workbook, "stock_workbook")
    macro_path = locate_extracted_file(extract_dir, args.macro_workbook, "macro_workbook")

    docx_text = read_docx_text(docx_path)
    keyword_results = extract_keywords(docx_text, DEFAULT_KEYWORDS)
    stock_sheets = read_excel_sheets(stock_path, "stock_workbook")
    macro_sheets = read_excel_sheets(macro_path, "macro_workbook")
    expected_macro_sheets = [sheet.strip() for sheet in args.expected_macro_sheets.split(",") if sheet.strip()]
    checks = validate_expectations(
        inventory=inventory,
        stock_sheets=stock_sheets,
        macro_sheets=macro_sheets,
        expected_stock_sheet_count=args.expected_stock_sheet_count,
        expected_macro_sheets=expected_macro_sheets,
    )

    write_inventory_json(
        output_dir / "project_inventory.json",
        args,
        extract_dir,
        inventory,
        docx_path,
        stock_path,
        macro_path,
        keyword_results,
        stock_sheets,
        macro_sheets,
        checks,
    )
    write_requirements_md(output_dir / "task_requirements.md", keyword_results, stock_sheets, macro_sheets, checks)
    write_assumption_ledger_csv(output_dir / "assumption_ledger_seed.csv")

    failed_checks = [check for check in checks if not check["passed"]]
    print(f"[OK] 已生成: {output_dir / 'project_inventory.json'}")
    print(f"[OK] 已生成: {output_dir / 'task_requirements.md'}")
    print(f"[OK] 已生成: {output_dir / 'assumption_ledger_seed.csv'}")
    if failed_checks:
        print("[WARN] 以下验收检查未通过:")
        for check in failed_checks:
            print(f"  - {check['name']}: expected={check['expected']} actual={check['actual']}")
        return 1 if args.fail_on_check else 0
    print("[OK] 所有默认验收检查通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
